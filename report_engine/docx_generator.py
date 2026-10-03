from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from config.constants import DATE_FMT
from config.settings import SIGNATORY_COMPANY
from report_engine import seal_inserter
from report_engine.formatting import CONTENT_PT, REPORT_FONT, apply_report_font, cell_padding
from report_engine.table_builder import build_parameter_table

BLUE = RGBColor(0x1F, 0x4E, 0x8C)
LINE_PT = 13.8  # height of one blank line of 12 pt Times New Roman
SECTION_GAP_PT = 3 * LINE_PT  # same gap below WATER REPORT and above the sample heading


def _bottom_border(paragraph, color="1F4E8C", sz=12) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", "1"), ("color", color)):
        el.set(qn(f"w:{k}"), v)
    b.append(el)
    pPr.append(b)


def _top_border(paragraph, color="1F4E8C", sz=8) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    el = OxmlElement("w:top")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", "4"), ("color", color)):
        el.set(qn(f"w:{k}"), v)
    b.append(el)
    pPr.append(b)


def _setup_page(doc) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(3.2), Cm(2.8)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)


def build_template(path: Path, footer_lines: List[str]) -> None:
    """Institutional template: shared letterhead header + office-specific footer address."""
    doc = Document()
    _setup_page(doc)
    sec = doc.sections[0]

    h = sec.header.paragraphs[0]
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("NIRMAAN")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(26), BLUE
    h2 = sec.header.add_paragraph()
    h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h2.add_run("WATERTECH SOLUTIONS  |  Engineering Water, Empowering Industries")
    r.font.size, r.font.color.rgb = Pt(8.5), BLUE
    _bottom_border(h2)

    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _top_border(f)
    for i, line in enumerate(footer_lines):
        if i:
            f.add_run().add_break(WD_BREAK.LINE)
        r = f.add_run(line)
        r.font.size, r.bold = Pt(8), (i == 0)
    # body must contain no placeholder text; generator appends below
    path.parent.mkdir(parents=True, exist_ok=True)
    apply_report_font(doc)
    doc.save(str(path))


def _fmt(d) -> str:
    return d.strftime(DATE_FMT) if hasattr(d, "strftime") else str(d or "")


def _info_table(doc, data: Dict) -> None:
    rows = [
        ("SAMPLE COLLECTION DATE", _fmt(data.get("collection_date"))),
        ("SAMPLE ANALYSIS DATE", _fmt(data.get("analysis_date"))),
        ("CLIENT NAME", data.get("client_name", "")),
        ("ADDRESS", data.get("client_address", "")),
        ("CONTACT PERSON", data.get("contact_person", "")),
    ]
    # Label | ":" | value columns keep every colon vertically aligned.
    t = doc.add_table(rows=0, cols=3)
    t.autofit = False
    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    label_w, colon_w = Cm(6.0), Cm(0.6)  # label fits "SAMPLE COLLECTION DATE" at 12 pt
    widths = (label_w, colon_w, width - label_w - colon_w)
    for column, col_width in zip(t.columns, widths):
        column.width = col_width
    table_width = t._tbl.tblPr.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(round(width / 635)))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        # Bold black rules above Sample Collection Date and below Contact Person.
        border.set(qn("w:val"), "single" if edge in ("top", "bottom") else "nil")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for index, (label, value) in enumerate(rows):
        cells = t.add_row().cells
        for col, (cell, text, col_width) in enumerate(zip(cells, (label, ":", value), widths)):
            cell.width = col_width
            cell_padding(cell, vertical=20)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 1 else WD_ALIGN_PARAGRAPH.LEFT
            # Only the 1 pt cell margins separate the rules from the content.
            # Major-section spacing belongs outside this compact table.
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = index < len(rows)-1
            run = paragraph.add_run(text)
            run.bold, run.font.size = False, Pt(CONTENT_PT)


def _page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def generate_docx(data: Dict, template_path: Path, seal_path: Path, out_path: Path) -> Path:
    """data keys: see services.report_service.save_report."""
    doc = Document(str(template_path))
    normal = doc.styles["Normal"]
    normal.font.name = REPORT_FONT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.0
    normal.font.size = Pt(CONTENT_PT)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    columns: List[str] = list(data["sample_columns"])
    params = data["parameters"]

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("WATER REPORT")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(16), BLUE
    t.paragraph_format.space_before = Pt(6)
    t.paragraph_format.space_after = Pt(SECTION_GAP_PT)
    t.paragraph_format.keep_with_next = True

    _info_table(doc, data)

    # Blank-line equivalents: ~3 lines after Contact Person, ~2 before "Dear Sir,".
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(SECTION_GAP_PT)
    title.paragraph_format.space_after = Pt(2 * LINE_PT)
    title.paragraph_format.keep_with_next = True
    r = title.add_run(" & ".join(columns).upper() + " ANALYSIS REPORT")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(14), BLUE

    salutation = doc.add_paragraph("Dear Sir,")
    salutation.paragraph_format.space_before = Pt(0)
    salutation.paragraph_format.space_after = Pt(LINE_PT)
    _keep(salutation)
    intro = doc.add_paragraph()
    client = data.get("client_name", "").strip() or "your organisation"
    intro.add_run(
        f"Please find below the water analysis results for {client}. "
        "The measured parameters for the submitted samples are presented for your review."
    ).font.size = Pt(CONTENT_PT)
    intro.paragraph_format.space_after = Pt(12)
    _keep(intro)

    # Let the document renderer paginate using actual row and text heights.
    # Table headers repeat and individual rows stay intact on each page.
    section = doc.sections[0]
    usable_cm = (section.page_width - section.left_margin - section.right_margin) / 360000
    build_parameter_table(doc, params, columns, total_width_cm=usable_cm)

    _remarks_block(doc, data, seal_path)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    apply_report_font(doc)
    doc.save(str(out_path))
    return out_path


def _keep(p) -> None:
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True


def _remarks_block(doc, data: Dict, seal_path: Path) -> None:
    head = doc.add_paragraph()
    head.paragraph_format.space_before = Pt(10)
    _keep(head)
    r = head.add_run("REMARKS")
    r.bold, r.font.size = True, Pt(CONTENT_PT)

    remark = data.get("remark_final") or data.get("remark_raw") or ""
    # Each supplied line is a complete point; keep sentences and decimal values intact.
    points = [re.sub(r"^\s*(?:[•*\-]\s*|\d+[.)]\s+)", "", line).strip()
              for line in remark.splitlines() if line.strip()]
    for point in filter(None, points):
        body = doc.add_paragraph()
        _keep(body)
        body.paragraph_format.left_indent = Cm(0.45)
        body.paragraph_format.first_line_indent = Cm(-0.35)
        body.paragraph_format.space_after = Pt(4)
        body.add_run("•  " + point).font.size = Pt(CONTENT_PT)

    for_p = doc.add_paragraph()
    for_p.paragraph_format.space_before = Pt(10)
    for_p.paragraph_format.space_after = Pt(3)
    _keep(for_p)
    for_p.add_run("For,").font.size = Pt(CONTENT_PT)
    co = doc.add_paragraph()
    co.paragraph_format.space_after = Pt(3)
    _keep(co)
    co.add_run(SIGNATORY_COMPANY).bold = True

    seal_p = doc.add_paragraph()
    seal_p.paragraph_format.space_after = Pt(3)
    _keep(seal_p)
    seal_inserter.insert_seal(seal_p, seal_path)

    name = doc.add_paragraph()
    name.paragraph_format.space_after = Pt(3)
    _keep(name)
    name.add_run(data.get("technician", "")).bold = True
    email = (data.get("technician_email") or "").strip()
    if not email:
        return
    mail = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), mail.part.relate_to(f"mailto:{email}", RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, value in (("color", "1F4E8C"), ("u", "single"), ("sz", str(CONTENT_PT * 2))):
        prop = OxmlElement(f"w:{tag}")
        prop.set(qn("w:val"), value)
        props.append(prop)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = email
    run.append(text)
    link.append(run)
    mail._p.append(link)
