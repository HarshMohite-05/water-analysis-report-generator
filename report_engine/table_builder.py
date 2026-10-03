from __future__ import annotations

from typing import Dict, List, Sequence

from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from report_engine.formatting import CONTENT_PT, REPORT_FONT, cell_padding, format_number

HEADER_FILL = "D9E6F2"


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _row_flags(row, header: bool) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def _write(cell, text: str, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=CONTENT_PT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = REPORT_FONT
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0, 0, 0)


def build_parameter_table(doc, params: Sequence[Dict], columns: List[str],
                          total_width_cm: float = 17.4):
    """One table: Parameter | Unit | <sample columns...>. params: [{name, unit, values{}}]."""
    n_cols = 2 + len(columns)
    table = doc.add_table(rows=1, cols=n_cols)
    if "Table Grid" in doc.styles:
        table.style = "Table Grid"
    else:
        # Imported office templates may omit Word's built-in table styles.
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            for name, value in (("val", "single"), ("sz", "4"), ("color", "000000")):
                border.set(qn(f"w:{name}"), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    name_w, unit_w = 5.2, 2.0
    sample_w = (total_width_cm - name_w - unit_w) / max(len(columns), 1)
    widths = [name_w, unit_w] + [sample_w] * len(columns)

    for column, width in zip(table.columns, widths):
        column.width = Cm(width)

    hdr = table.rows[0]
    _row_flags(hdr, header=True)
    for i, label in enumerate(["PARAMETER", "UNIT"] + [c.upper() for c in columns]):
        _write(hdr.cells[i], label, bold=True)
        _shade(hdr.cells[i], HEADER_FILL)

    for p in params:
        row = table.add_row()
        _row_flags(row, header=False)
        _write(row.cells[0], p["name"], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        _write(row.cells[1], p.get("unit", ""))
        for i, col in enumerate(columns):
            _write(row.cells[2 + i], format_number((p.get("values") or {}).get(col, "")))

    for row in table.rows:
        row.height = Cm(0.77)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)
            cell_padding(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table
