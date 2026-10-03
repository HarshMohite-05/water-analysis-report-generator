import subprocess

import pytest
from docx import Document
from docx.oxml.ns import qn

from report_engine import pdf_converter
from report_engine.table_builder import build_parameter_table


def test_parameter_table_supports_template_without_grid_style(tmp_path):
    doc = Document()
    doc.styles["Table Grid"].delete()
    build_parameter_table(doc, [{"name": "pH", "unit": "", "values": {"Feed": "7"}}], ["Feed"])
    path = tmp_path / "report.docx"
    doc.save(path)
    table = Document(path).tables[0]
    assert table.cell(1, 2).text == "7"
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert len(borders) == 6


def test_conversion_timeout_is_actionable(monkeypatch, tmp_path):
    monkeypatch.setattr(pdf_converter, "_soffice", lambda: "/test/soffice")

    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("soffice", 120)

    monkeypatch.setattr(pdf_converter.subprocess, "run", timeout)
    with pytest.raises(pdf_converter.PdfConversionError, match="LibreOffice could not convert"):
        pdf_converter.docx_to_pdf(tmp_path / "report.docx", tmp_path)
