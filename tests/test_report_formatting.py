from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm
from report_engine.docx_generator import generate_docx
from services.template_service import template_paths
from tests.test_reports import _payload


def test_requested_report_formatting(tmp_path):
    payload = _payload(6)
    payload['remark_final'] = 'TDS is 2.5 mg/L. Review treatment.\n- Check the filter.'
    tpl, seal = template_paths(payload['office'])
    doc = Document(generate_docx(payload, tpl, seal, tmp_path / 'report.docx'))
    title = next(p for p in doc.paragraphs if p.text == 'WATER REPORT')
    assert title.runs[0].bold and str(title.runs[0].font.color.rgb) == '1F4E8C'
    assert any(p.text == 'Dear Sir,' for p in doc.paragraphs)
    assert any(payload['client_name'] in p.text for p in doc.paragraphs)
    assert not any('Analysed by:' in p.text for p in doc.paragraphs)
    bullets = [p.text for p in doc.paragraphs if p.text.startswith('•')]
    assert bullets == ['•  TDS is 2.5 mg/L. Review treatment.', '•  Check the filter.']
    assert any(r.reltype == RT.HYPERLINK and r.target_ref == 'mailto:technician@example.test'
               for r in doc.part.rels.values())
    assert doc.tables[0].autofit is False
    info = doc.tables[0]
    assert len(info.columns) == 3  # label | ':' | value, so colons line up
    assert [c.text for c in info.rows[2].cells] == ['CLIENT NAME', ':', 'Shree Prakash Textiles']
    assert {row.cells[1].text for row in info.rows} == {':'}
    assert not any(r.bold for c in info._cells for par in c.paragraphs for r in par.runs)
    for row in doc.tables[1].rows:
        assert abs(row.height - Cm(.77)) < 1000
        assert row._tr.trPr.find(qn('w:cantSplit')) is not None


def test_empty_remarks_do_not_create_bullets(tmp_path):
    payload = _payload(1)
    tpl, seal = template_paths(payload['office'])
    for index, remark in enumerate(['', ' \n\n ', '-\n•\n*\n1. ']):
        payload['remark_raw'] = ''
        payload['remark_final'] = remark
        doc = Document(generate_docx(payload, tpl, seal, tmp_path / f'empty{index}.docx'))
        assert not any(p.text.startswith('•') for p in doc.paragraphs)


def test_old_payload_without_received_date_still_loads():
    from services.report_service import _de, _ser
    payload = _payload(1)
    assert _de(_ser(payload)) == payload


def test_fixed_typography_alignment_padding_and_seal(tmp_path):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    for office in ('Mumbai Office', 'Ahmedabad Office'):
        data = _payload(6)
        data['office'] = office
        tpl, seal = template_paths(office)
        doc = Document(generate_docx(data, tpl, seal, tmp_path / f'{office}.docx'))
        headings = [p for p in doc.paragraphs if 'REPORT' in p.text]
        assert [p.runs[0].font.size.pt for p in headings] == [16, 14]
        for heading in headings:
            assert heading.runs[0].bold
            assert str(heading.runs[0].font.color.rgb) == '1F4E8C'
            assert heading.alignment == 1
        for run in doc.element.iter(qn('w:r')):
            fonts = run.find(qn('w:rPr')).find(qn('w:rFonts'))
            assert all(fonts.get(qn('w:' + slot)) == 'Times New Roman'
                       for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'))
        info = doc.tables[0]
        assert len({row.cells[0].width for row in info.rows}) == 1
        borders = info._tbl.tblPr.find(qn('w:tblBorders'))
        assert all(borders.find(qn('w:' + edge)).get(qn('w:val')) == 'single'
                   for edge in ('top', 'bottom'))
        for row in doc.tables[1].rows:
            for cell in row.cells:
                assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
                padding = cell._tc.tcPr.find(qn('w:tcMar'))
                assert padding.find(qn('w:top')).get(qn('w:w')) == padding.find(qn('w:bottom')).get(qn('w:w'))
        assert abs(doc.inline_shapes[0].width - Cm(3.2 * .82)) < 100


def test_format_number_adds_thousands_separators():
    from report_engine.formatting import format_number
    cases = {'10000': '10,000', '1234567': '1,234,567', '10000.50': '10,000.50',
             '999': '999', '7.2': '7.2', '0.005': '0.005', '-12500': '-12,500',
             ' 25000 ': '25,000', '10,000': '10,000', '': '', 'BDL': 'BDL',
             '<0.1': '<0.1', '6.5-8.5': '6.5-8.5', 'Nil': 'Nil'}
    for raw, expected in cases.items():
        assert format_number(raw) == expected, raw


def test_report_table_shows_thousands_separators(tmp_path):
    payload = _payload(1)
    payload['parameters'][0]['values'] = {'RO Feed Water': '10000', 'RO 3rd Stage Reject Water': '7.25'}
    tpl, seal = template_paths(payload['office'])
    doc = Document(generate_docx(payload, tpl, seal, tmp_path / 'numbers.docx'))
    assert [c.text for c in doc.tables[1].rows[1].cells[2:]] == ['10,000', '7.25']


def test_blank_line_spacing_around_heading_and_greeting(tmp_path):
    from report_engine.docx_generator import LINE_PT
    payload = _payload(1)
    tpl, seal = template_paths(payload['office'])
    doc = Document(generate_docx(payload, tpl, seal, tmp_path / 'spacing.docx'))
    texts = [p.text for p in doc.paragraphs]
    heading = doc.paragraphs[texts.index('RO FEED WATER & RO 3RD STAGE REJECT WATER ANALYSIS REPORT')]
    greeting = doc.paragraphs[texts.index('Dear Sir,')]
    from pytest import approx
    assert heading.paragraph_format.space_before.pt == approx(3 * LINE_PT, abs=.05)  # after Contact Person
    assert heading.paragraph_format.space_after.pt == approx(2 * LINE_PT, abs=.05)   # before Dear Sir,
    assert greeting.paragraph_format.space_after.pt == approx(LINE_PT, abs=.05)


def test_content_is_12pt_and_info_rules_are_bold_black(tmp_path):
    payload = _payload(3)
    tpl, seal = template_paths(payload['office'])
    doc = Document(generate_docx(payload, tpl, seal, tmp_path / 'size.docx'))
    headings = {'WATER REPORT', 'RO FEED WATER & RO 3RD STAGE REJECT WATER ANALYSIS REPORT'}
    sized = [r for p in doc.paragraphs if p.text not in headings for r in p.runs if r.font.size]
    sized += [r for t in doc.tables for c in t._cells for p in c.paragraphs for r in p.runs if r.text]
    assert sized and all(r.font.size.pt == 12 for r in sized)  # unsized runs inherit 12 pt Normal
    assert doc.styles['Normal'].font.size.pt == 12
    borders = doc.tables[0]._tbl.tblPr.find(qn('w:tblBorders'))
    for edge in ('top', 'bottom'):  # above Sample Collection Date, below Contact Person
        rule = borders.find(qn('w:' + edge))
        assert (rule.get(qn('w:val')), rule.get(qn('w:sz')), rule.get(qn('w:color'))) == ('single', '12', '000000')


def test_gap_below_title_matches_gap_above_sample_heading(tmp_path):
    import fitz
    from report_engine import pdf_converter
    payload = _payload(2)
    tpl, seal = template_paths(payload['office'])
    pdf = pdf_converter.docx_to_pdf(generate_docx(payload, tpl, seal, tmp_path / 'gap.docx'), tmp_path)
    page = fitz.open(pdf)[0]
    title = page.search_for('WATER REPORT')[0]
    heading = page.search_for('RO FEED WATER & RO')[0]
    rules = sorted(d['rect'].y0 for d in page.get_drawings()
                   if d['rect'].width > 300 and d['rect'].height < 3 and title.y1 < d['rect'].y0 < heading.y0)
    assert abs((rules[0] - title.y1) - (heading.y0 - rules[-1])) < 2  # points


def test_no_remarks_omits_heading_and_text_but_keeps_signature(tmp_path):
    import fitz
    from report_engine.pdf_converter import docx_to_pdf
    data = _payload(1)
    data.update(include_remarks=False, remark_raw='Hidden raw remark', remark_final='Hidden final remark')
    tpl, seal = template_paths(data['office'])
    path = generate_docx(data, tpl, seal, tmp_path / 'no_remarks.docx')
    doc = Document(path)
    assert len(doc.inline_shapes) == 1
    with fitz.open(docx_to_pdf(path, tmp_path)) as pdf:
        pdf_text = '\n'.join(p.get_text() for p in pdf)
    for text in ('\n'.join(p.text for p in doc.paragraphs), pdf_text):
        assert 'REMARKS' not in text
        assert 'Hidden raw remark' not in text
        assert 'Hidden final remark' not in text
        assert data['technician'] in text
