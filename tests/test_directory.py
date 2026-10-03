from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from database.database import session_scope
from database.models import User
from services import auth_service, dataset_service, report_service


def change_employee(employee_id, column, value):
    book = load_workbook(dataset_service.EMPLOYEE_FILE)
    sheet = book.active
    row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value == employee_id)
    sheet.cell(row, column).value = value
    book.save(dataset_service.EMPLOYEE_FILE)
    book.close()


@pytest.mark.parametrize('identifier,password,role,allowed', [
    ('ADM001', 'ChangeMe#1', 'admin', True),
    ('DIR1', 'DirectorTest!', 'director', True),
    ('TECH1', 'TechTest!', 'technician', False),
    ('SALE1', 'SalesTest!', 'sales executive', False),
])
def test_roles_are_checked_on_admin_page_and_writes(isolated_reports, identifier, password, role, allowed):
    user = auth_service.authenticate(identifier, password)
    assert user['role'] == role
    assert 'password' not in user and 'password_hash' not in user
    app = AppTest.from_file(str(Path('pages/05_Admin_Panel.py').resolve()))
    app.session_state['user'] = user
    app.run()
    assert not app.exception
    if allowed:
        assert app.tabs and not app.error
        dataset_service.require_admin(user['id'])
    else:
        assert app.error and not app.tabs
        original = dataset_service.CLIENT_FILE.read_bytes()
        with pytest.raises(PermissionError):
            dataset_service.save_client(user['id'], {})
        with pytest.raises(PermissionError):
            dataset_service.save_employee(user['id'], {})
        assert dataset_service.CLIENT_FILE.read_bytes() == original


def test_workbook_updates_refresh_authentication_and_client_links(isolated_reports):
    user = auth_service.authenticate('tech@example.test', 'TechTest!')
    assert user and user['employee_id'] == 'TECH1'
    clients = report_service.list_clients()
    assert all(c['assigned_employee_id'] == 'TECH1' for c in clients)
    assert all(c['assigned_employee_email'] == 'tech@example.test' for c in clients)
    change_employee('TECH1', 6, 'UpdatedTest!')
    assert auth_service.authenticate('TECH1', 'TechTest!') is None
    assert auth_service.authenticate('TECH1', 'UpdatedTest!')['id'] == user['id']
    assert auth_service.authenticate('%', 'ChangeMe#1') is None
    assert auth_service.authenticate('EMP001', 'ChangeMe#1') is None


def test_admin_changes_write_back_and_do_not_expose_passwords(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    tech = next(e for e in dataset_service.employee_directory() if e['employee_id'] == 'TECH1')
    assert 'password' not in tech and 'password_hash' not in tech
    tech['role'] = 'sales executive'
    tech['password'] = 'ReplacementTest!'
    dataset_service.save_employee(admin['id'], tech)
    assert auth_service.authenticate('TECH1', 'ReplacementTest!')['role'] == 'sales executive'
    assert 'pbkdf2$' in load_workbook(dataset_service.EMPLOYEE_FILE).active.cell(4, 6).value
    dataset_service.save_client(admin['id'], dict(name='Mumbai Client', office='Ahmedabad Office',
        address='Updated address', contact_person='New contact', email='updated@example.test', assigned_employee_id='DIR1'))
    client = next(c for c in report_service.list_clients() if c['name'] == 'Mumbai Client')
    assert client['assigned_employee_id'] == 'DIR1' and client['office'] == 'Ahmedabad Office'
    assert client['address'] == 'Updated address'
    assert list((isolated_reports / 'storage/dataset_backups').glob('*.xlsx'))


def test_role_revocation_and_logout_clear_access(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    change_employee('ADM001', 3, 'Technician')
    with pytest.raises(PermissionError):
        dataset_service.require_admin(admin['id'])
    app = AppTest.from_file(str(Path('pages/05_Admin_Panel.py').resolve()))
    app.session_state['user'] = admin
    app.run()
    assert app.error and not app.tabs
    source = '''
import streamlit as st
from components.authentication import logout
st.session_state['user'] = {'id': 1, 'role': 'admin'}
st.session_state['draft'] = {'values': {'pH': '7'}}
logout()
assert len(st.session_state) == 0
'''
    assert not AppTest.from_string(source).run().exception


def test_invalid_dataset_does_not_partially_update_database(isolated_reports):
    before = dataset_service.employee_directory()
    change_employee('TECH1', 3, '')
    with pytest.raises(dataset_service.DatasetError):
        dataset_service.sync_directory()
    with session_scope() as session:
        assert session.query(User).filter_by(employee_id='TECH1').one().role == 'technician'
    assert len(before) == 4


def test_sync_preserves_ids_hashes_and_retired_history(isolated_reports):
    first = auth_service.authenticate('TECH1', 'TechTest!')
    with session_scope() as session:
        original_hash = session.get(User, first['id']).password_hash
    dataset_service.sync_directory(force=True)
    with session_scope() as session:
        assert session.get(User, first['id']).password_hash == original_hash
    book = load_workbook(dataset_service.EMPLOYEE_FILE)
    book.active.delete_rows(4)
    book.save(dataset_service.EMPLOYEE_FILE)
    assert auth_service.authenticate('TECH1', 'TechTest!') is None
    with session_scope() as session:
        assert session.get(User, first['id']) is not None
        assert not session.get(User, first['id']).active


def test_add_employee_and_client_without_hardcoded_directory(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    dataset_service.save_employee(admin['id'], dict(employee_id='NEW1', full_name='New Employee',
        role='technician', contact_number='999', email='new@example.test', password='NewTest!'))
    new = auth_service.authenticate('NEW1', 'NewTest!')
    assert new and new['full_name'] == 'New Employee'
    dataset_service.save_client(admin['id'], dict(name='New Client', office='Mumbai Office',
        address='New address', contact_person='New contact', email='client2@example.test', assigned_employee_id='NEW1'))
    client = next(c for c in report_service.list_clients() if c['name'] == 'New Client')
    assert client['assigned_employee_id'] == 'NEW1'
    assert client['assigned_employee_email'] == 'new@example.test'
    assert client['assigned_technician'] == 'New Employee'


def test_invalid_admin_edit_leaves_workbook_unchanged(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    original = dataset_service.EMPLOYEE_FILE.read_bytes()
    with pytest.raises(dataset_service.DatasetError):
        dataset_service.save_employee(admin['id'], dict(employee_id='NEW1', full_name='New Employee',
            role='technician', contact_number='999', email='admin@example.test', password='NewTest!'))
    assert dataset_service.EMPLOYEE_FILE.read_bytes() == original


def test_employee_rename_preserves_id_and_updates_client_workbook(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    before = auth_service.authenticate('TECH1', 'TechTest!')
    values = next(e for e in dataset_service.employee_directory() if e['employee_id'] == 'TECH1')
    values['full_name'] = 'Renamed Technician'
    dataset_service.save_employee(admin['id'], values)
    after = auth_service.authenticate('TECH1', 'TechTest!')
    assert after['id'] == before['id']
    assert after['full_name'] == 'Renamed Technician'
    assert all(c['assigned_employee_id'] == 'TECH1' and c['assigned_technician'] == 'Renamed Technician'
               for c in report_service.list_clients())
    book = load_workbook(dataset_service.CLIENT_FILE)
    assert book.active.cell(2, 6).value == 'Renamed Technician'


def test_client_rename_updates_existing_database_record(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    before = next(c for c in report_service.list_clients() if c['name'] == 'Mumbai Client')
    dataset_service.save_client(admin['id'], dict(name='Renamed Client', office='Mumbai Office',
        address='Updated address', contact_person='Updated contact', email='new@example.test',
        assigned_employee_id='TECH1'), original_name='Mumbai Client')
    clients = report_service.list_clients()
    assert len(clients) == 2
    after = next(c for c in clients if c['name'] == 'Renamed Client')
    assert after['id'] == before['id'] and after['address'] == 'Updated address'
    assert not any(c['name'] == 'Mumbai Client' for c in clients)


def test_linked_rename_rolls_back_both_workbooks_on_sync_failure(isolated_reports, monkeypatch):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    values = next(e for e in dataset_service.employee_directory() if e['employee_id'] == 'TECH1')
    originals = {p: p.read_bytes() for p in (dataset_service.EMPLOYEE_FILE, dataset_service.CLIENT_FILE)}
    sync = dataset_service.sync_directory
    def failing_sync(force=False, **kwargs):
        if force:
            raise RuntimeError('Simulated database failure')
        return sync(force=force, **kwargs)
    monkeypatch.setattr(dataset_service, 'sync_directory', failing_sync)
    values['full_name'] = 'Renamed Technician'
    with pytest.raises(RuntimeError):
        dataset_service.save_employee(admin['id'], values)
    assert all(p.read_bytes() == original for p, original in originals.items())


def test_admin_can_edit_existing_names_in_panel(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    app = AppTest.from_file(str(Path('pages/05_Admin_Panel.py').resolve()))
    app.session_state['user'] = admin
    app.run()
    def field(kind, label):
        return next(w for w in getattr(app, kind) if w.label == label)
    field('selectbox', 'Employee to manage').select('TECH1').run()
    field('text_input', 'Employee name').input('Edited in Admin Panel')
    field('button', 'Save employee').click().run()
    assert not app.exception and not app.error
    assert auth_service.authenticate('TECH1', 'TechTest!')['full_name'] == 'Edited in Admin Panel'
    field('selectbox', 'Client to manage').select('Mumbai Client').run()
    field('text_input', 'Client name').input('Client Edited in Admin Panel')
    field('button', 'Save client').click().run()
    assert not app.exception and not app.error
    assert any(c['name'] == 'Client Edited in Admin Panel' for c in report_service.list_clients())
    assert 'Reports' in [tab.label for tab in app.tabs]


def test_delete_employee_disables_login_and_preserves_history(isolated_reports):
    from database.models import Report
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    sales = auth_service.authenticate('SALE1', 'SalesTest!')
    with session_scope() as session:
        report = Report(office='Mumbai Office', client_name='Historic client', created_by=sales['id'], data={})
        session.add(report)
        session.flush()
        report_id = report.id
    dataset_service.delete_employee(admin['id'], 'SALE1')
    assert auth_service.authenticate('SALE1', 'SalesTest!') is None
    assert all(e['employee_id'] != 'SALE1' for e in dataset_service.read_directory()[0])
    dataset_service.sync_directory(force=True)
    with session_scope() as session:
        assert session.get(User, sales['id']).active is False
        assert session.get(Report, report_id).created_by == sales['id']
    assert list((isolated_reports / 'storage' / 'dataset_backups').glob('*.xlsx'))


def test_delete_guards_and_non_admin_denial(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    tech = auth_service.authenticate('TECH1', 'TechTest!')
    before = dataset_service.EMPLOYEE_FILE.read_bytes()
    with pytest.raises(dataset_service.DatasetError, match='own account'):
        dataset_service.delete_employee(admin['id'], 'ADM001')
    with pytest.raises(dataset_service.DatasetError, match='Reassign'):
        dataset_service.delete_employee(admin['id'], 'TECH1')
    with pytest.raises(PermissionError):
        dataset_service.delete_employee(tech['id'], 'SALE1')
    with pytest.raises(PermissionError):
        dataset_service.delete_client(tech['id'], 'Mumbai Client')
    assert dataset_service.EMPLOYEE_FILE.read_bytes() == before


def test_delete_client_rolls_back_on_failure(isolated_reports, monkeypatch):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    before = dataset_service.CLIENT_FILE.read_bytes()
    original_sync = dataset_service.sync_directory
    def fail_on_write(*args, **kwargs):
        if kwargs.get('force'):
            raise RuntimeError('test sync failure')
        return original_sync(*args, **kwargs)
    monkeypatch.setattr(dataset_service, 'sync_directory', fail_on_write)
    with pytest.raises(RuntimeError):
        dataset_service.delete_client(admin['id'], 'Mumbai Client')
    assert dataset_service.CLIENT_FILE.read_bytes() == before


def test_admin_delete_buttons_remove_selected_records(isolated_reports):
    app = AppTest.from_file(str(Path('pages/05_Admin_Panel.py').resolve()))
    app.session_state['user'] = auth_service.authenticate('ADM001', 'ChangeMe#1')
    app.run()
    def field(kind, label):
        return next(w for w in getattr(app, kind) if w.label == label)
    field('selectbox', 'Employee to manage').select('SALE1').run()
    assert field('button', 'Delete employee').disabled
    field('checkbox', 'Confirm deletion of Sales Person').check().run()
    field('button', 'Delete employee').click().run()
    assert not app.exception and not app.error
    assert all(e['employee_id'] != 'SALE1' for e in dataset_service.employee_directory())
    # Reload after deletion: AppTest retains widgets removed by st.rerun.
    app = AppTest.from_file(str(Path('pages/05_Admin_Panel.py').resolve()))
    app.session_state['user'] = auth_service.authenticate('ADM001', 'ChangeMe#1')
    app.run()
    field('selectbox', 'Client to manage').select('Mumbai Client').run()
    field('checkbox', 'Confirm deletion of Mumbai Client').check().run()
    field('button', 'Delete client').click().run()
    assert not app.exception and not app.error
    assert all(c['name'] != 'Mumbai Client' for c in report_service.list_clients())
    from database.models import Client
    with session_scope() as session:
        assert session.query(Client).filter_by(name='Mumbai Client').one().active is False


def test_custom_role_can_be_created_and_edited_in_admin_panel(isolated_reports):
    app = AppTest.from_file(str(Path('pages/05_Admin_Panel.py').resolve()))
    app.session_state['user'] = auth_service.authenticate('ADM001', 'ChangeMe#1')
    app.run()
    def field(kind, label):
        return next(w for w in getattr(app, kind) if w.label == label)
    assert field('selectbox', 'Employee to manage').value == 'New employee'
    assert field('selectbox', 'Client to manage').value == 'New client'
    for label, value in [('Employee ID', 'INT1'), ('Employee name', 'New Intern'),
                         ('Employee email', 'intern@example.test'), ('Role', 'Research Intern'),
                         ('Password (leave blank to keep existing)', 'InternTest!')]:
        field('text_input', label).input(value)
    field('button', 'Save employee').click().run()
    assert not app.exception and not app.error
    intern = auth_service.authenticate('INT1', 'InternTest!')
    assert intern['role'] == 'research intern'
    with pytest.raises(PermissionError):
        dataset_service.require_admin(intern['id'])
    field('selectbox', 'Employee to manage').select('INT1').run()
    assert field('text_input', 'Role').value == 'research intern'
    field('text_input', 'Role').input('Water Quality Laboratory Assistant')
    field('button', 'Save employee').click().run()
    assert not app.exception and not app.error
    assert auth_service.authenticate('INT1', 'InternTest!')['role'] == 'water quality laboratory assistant'


def test_blank_role_is_rejected_without_changing_directory(isolated_reports):
    admin = auth_service.authenticate('ADM001', 'ChangeMe#1')
    values = next(e for e in dataset_service.employee_directory() if e['employee_id'] == 'TECH1')
    before = dataset_service.EMPLOYEE_FILE.read_bytes()
    values['role'] = ' '
    with pytest.raises(dataset_service.DatasetError, match='must have a role'):
        dataset_service.save_employee(admin['id'], values)
    assert dataset_service.EMPLOYEE_FILE.read_bytes() == before
