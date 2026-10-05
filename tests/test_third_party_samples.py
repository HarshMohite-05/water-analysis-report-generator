from pathlib import Path

import fitz
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from report_engine.docx_generator import generate_docx
from report_engine.pdf_converter import docx_to_pdf
from services import report_service, template_service
from tests.test_reports import _payload
from tests.test_workflow import widget, run_ui


def test_anonymous_validation_preserves_other_requirements():
    payload = _payload(1)
    payload['client_name'] = ''
    assert 'Select a client.' in report_service.validate(payload)
    payload['third_party_sample'] = True
    assert report_service.validate(payload) == []
    payload['parameters'] = []
    assert 'Select at least one analysis parameter.' in report_service.validate(payload)


@pytest.mark.parametrize('office', ['Mumbai Office', 'Ahmedabad Office'])
def test_anonymous_documents_exclude_even_stale_client_details(tmp_path, office):
    payload = _payload(2)
    payload.update(third_party_sample=True, office=office)
    tpl, seal = template_service.template_paths(office)
    path = generate_docx(payload, tpl, seal, tmp_path / 'anonymous.docx')
    doc = Document(path)
    assert len(doc.tables[0].rows) == 2
    text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for c in t._cells)
    with fitz.open(docx_to_pdf(path, tmp_path)) as pdf:
        pdf_text = '\n'.join(p.get_text() for p in pdf)
    for content in (text, pdf_text):
        for forbidden in ('CLIENT NAME', 'CONTACT PERSON', 'ADDRESS', payload['client_name'], payload['client_address'], payload['contact_person']):
            assert forbidden not in content
        assert 'SAMPLE COLLECTION DATE' in content
        assert 'third-party samples' in content


def test_anonymous_ui_save_reopen_and_return_to_normal(isolated_reports):
    app = AppTest.from_file(str(Path('app.py').resolve())).run()
    widget(app, 'text_input', 'Employee ID / Email').input('TECH1')
    widget(app, 'text_input', 'Password').input('TechTest!')
    widget(app, 'button', 'Login').click().run()
    app.switch_page('pages/03_Create_Report.py').run()
    widget(app, 'selectbox', 'Client Name').select('Mumbai Client')
    run_ui(app)
    widget(app, 'radio', 'Client details').set_value('N/A — Third-Party Sample Testing')
    run_ui(app)
    assert not any(w.label == 'Client Name' for w in app.selectbox)
    assert not any(w.label == 'Contact Person' for w in app.text_input)
    assert not any(w.label == 'Address' for w in app.text_area)
    assert not widget(app, 'button', '💾 Save draft').disabled
    widget(app, 'radio', 'Include remarks?').set_value('No')
    run_ui(app)
    widget(app, 'button', '💾 Save draft').click()
    run_ui(app)
    rep = report_service.get_report(app.session_state['draft']['report_id'])
    assert rep['data']['third_party_sample'] is True
    assert rep['data']['include_remarks'] is False
    assert rep['data']['remark_raw'] == rep['data']['remark_final'] == ''
    for field in ('client_name', 'client_address', 'contact_person', 'client_email', 'assigned_technician'):
        assert rep['data'][field] == ''
    assert rep['data']['assigned_employee_id'] is None
    app.session_state['open_report_id'] = rep['id']
    run_ui(app)
    assert widget(app, 'radio', 'Include remarks?').value == 'No'
    assert widget(app, 'radio', 'Client details').value == 'N/A — Third-Party Sample Testing'
    widget(app, 'radio', 'Client details').set_value('Available')
    run_ui(app)
    assert widget(app, 'button', '💾 Save draft').disabled
    widget(app, 'selectbox', 'Client Name').select('Mumbai Client')
    run_ui(app)
    assert not widget(app, 'button', '💾 Save draft').disabled
    assert widget(app, 'text_input', 'Contact Person').value == 'Contact'
