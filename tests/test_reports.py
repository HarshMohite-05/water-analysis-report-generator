from datetime import date

from docx import Document

from config.settings import TEMP_DIR
from report_engine import docx_generator
from services import template_service


def _payload(n_params=9):
    return dict(
        office="Ahmedabad Office", client_name="Shree Prakash Textiles", client_address="Addr",
        contact_person="Mr. R", collection_date=date(2026, 6, 9), analysis_date=date(2026, 6, 12),
        technician="Rajesh Kumar", technician_email="technician@example.test", sample_columns=["RO Feed Water", "RO 3rd Stage Reject Water"],
        parameters=[dict(name=f"P{i}", unit="Mg/L", values={"RO Feed Water": "1", "RO 3rd Stage Reject Water": "2"})
                    for i in range(n_params)],
        remark_raw="x", remark_final="The RO reject water shows higher TDS.",
    )


def test_docx_contains_all_rows_and_remark():
    template_service.ensure_templates()
    tpl, seal = template_service.template_paths("Ahmedabad Office")
    out = docx_generator.generate_docx(_payload(40), tpl, seal, TEMP_DIR / "t.docx")
    doc = Document(str(out))
    assert sum(len(t.rows) - 1 for t in doc.tables[1:]) == 40   # table 0 is the info block
    assert any("higher TDS" in p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 1  # seal