"""Excel-backed employee/client directory, synchronized into the existing database."""
from __future__ import annotations

import hashlib
import os
import re
import threading
from pathlib import Path
from tempfile import NamedTemporaryFile
from datetime import datetime
from zipfile import BadZipFile
from openpyxl.utils.exceptions import InvalidFileException

from openpyxl import load_workbook

from config.settings import BASE_DIR, OFFICES
from database.database import session_scope
from database.models import User, Client

EMPLOYEE_FILE = BASE_DIR / 'database' / 'Employee Data.xlsx'
CLIENT_FILE = BASE_DIR / 'database' / 'Client Data.xlsx'
ADMIN_ROLES = ('admin', 'director')
_LOCK = threading.RLock()
_last_signature = None


class DatasetError(ValueError):
    pass


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _key(value):
    return re.sub(r'[^a-z0-9]', '', _text(value).lower())


EMP_HEADERS = {'employeeid': 'employee_id', 'employeename': 'full_name', 'role': 'role',
               'contactnumber': 'contact_number', 'emailid': 'email', 'password': 'password'}
CLIENT_HEADERS = {'srno': 'serial', 'clientname': 'name', 'office': 'office', 'address': 'address',
                  'contactperson': 'contact_person', 'technician': 'assigned_technician',
                  'technicin': 'assigned_technician', 'mailid': 'email'}


def _sheet(workbook, headers, required):
    for sheet in workbook:
        for row in sheet.iter_rows(min_row=1, max_row=min(20, sheet.max_row)):
            columns = {headers[_key(c.value)]: c.column for c in row if _key(c.value) in headers}
            if set(required) <= columns.keys():
                return sheet, row[0].row, columns
    raise DatasetError('Dataset headers are missing or unrecognized.')


def _read(path, headers, required):
    try:
        workbook = load_workbook(path, data_only=True)
        sheet, header_row, columns = _sheet(workbook, headers, required)
        rows = []
        for number in range(header_row + 1, sheet.max_row + 1):
            row = {field: _text(sheet.cell(number, col).value) for field, col in columns.items()}
            if 'password' in columns:
                value = sheet.cell(number, columns['password']).value
                if isinstance(value, str):
                    row['password'] = value
            if any(row.values()):
                rows.append(row)
        workbook.close()
        return rows
    except (OSError, KeyError, BadZipFile, InvalidFileException) as exc:
        raise DatasetError(f'Cannot read {Path(path).name}; check that the workbook is available.') from exc


def read_directory(employee_file=None, client_file=None):
    employees = _read(employee_file or EMPLOYEE_FILE, EMP_HEADERS, EMP_HEADERS.values())
    clients = _read(client_file or CLIENT_FILE, CLIENT_HEADERS,
                    ('name', 'office', 'address', 'contact_person', 'assigned_technician', 'email'))
    ids, emails, names = set(), set(), {}
    for employee in employees:
        role = employee['role'].casefold().replace('_', ' ')
        employee['role'] = {'technicion': 'technician'}.get(role, role)
        employee['email'] = employee['email'].lower()
        if not all(employee[f] for f in ('employee_id', 'full_name', 'email', 'password')):
            raise DatasetError('Every employee must have an ID, name, email and password.')
        if not employee['role'].strip():
            raise DatasetError('Every employee must have a role.')
        identifier = employee['employee_id'].casefold()
        if identifier in ids or employee['email'] in emails:
            raise DatasetError('Employee IDs and email addresses must be unique.')
        ids.add(identifier); emails.add(employee['email'])
        names.setdefault(_key(employee['full_name']), []).append(employee)
    if not employees:
        raise DatasetError('Employee dataset is empty.')
    seen_clients = set()
    for client in clients:
        if not client['name'] or _key(client['name']) in seen_clients:
            raise DatasetError('Client names must be present and unique.')
        seen_clients.add(_key(client['name']))
        office = next((o for o in OFFICES if _key(client['office']) in (_key(o), _key(o.replace(' Office', '')))), None)
        if not office:
            raise DatasetError(f'Unknown office for client {client["name"]}.')
        client['office'] = office
        matches = names.get(_key(client['assigned_technician']), [])
        client['assigned_employee_id'] = matches[0]['employee_id'] if len(matches) == 1 else None
    return employees, clients


def sync_directory(force=False, client_rename=None):
    global _last_signature
    from database import database
    from services.auth_service import hash_password, verify_password
    with _LOCK:
        try:
            signature = (id(database.engine), str(EMPLOYEE_FILE), str(CLIENT_FILE),
                         hashlib.sha256(Path(EMPLOYEE_FILE).read_bytes()).digest(),
                         hashlib.sha256(Path(CLIENT_FILE).read_bytes()).digest())
        except OSError as exc:
            raise DatasetError('Employee Data.xlsx and Client Data.xlsx must be available in database/.') from exc
        if not force and signature == _last_signature:
            return
        employees, clients = read_directory()
        with session_scope() as session:
            incoming_ids = {e['employee_id'] for e in employees}
            incoming_emails = {e['email'] for e in employees}
            email_by_id = {e['employee_id']: e['email'] for e in employees}
            for user in session.query(User):
                user.active = False
                if user.email in incoming_emails and email_by_id.get(user.employee_id) != user.email:
                    user.email = f'archived-{user.id}@invalid.local'
            session.flush()
            for employee in employees:
                user = session.query(User).filter(User.employee_id == employee['employee_id']).first()
                if user is None:
                    user = User(employee_id=employee['employee_id'])
                    session.add(user)
                password = employee['password']
                if password.startswith('pbkdf2$'):
                    user.password_hash = password
                elif not user.password_hash or not verify_password(password, user.password_hash):
                    user.password_hash = hash_password(password)
                for field in ('full_name', 'email', 'role', 'contact_number'):
                    setattr(user, field, employee[field])
                user.active = True
                offices = {c['office'] for c in clients if c['assigned_employee_id'] == employee['employee_id']}
                if len(offices) == 1:
                    user.default_office = next(iter(offices))
            session.flush()
            if client_rename:
                old_name, new_name = client_rename
                existing = session.query(Client).filter(Client.name == old_name).first()
                if existing:
                    existing.name = new_name
                    session.flush()
            for client in session.query(Client):
                client.active = False
            for row in clients:
                client = session.query(Client).filter(Client.name == row['name']).first()
                if client is None:
                    client = Client(name=row['name'])
                    session.add(client)
                for field in ('office', 'address', 'contact_person', 'email', 'assigned_technician', 'assigned_employee_id'):
                    setattr(client, field, row[field])
                client.active = True
        _last_signature = signature


def employee_directory():
    sync_directory()
    with session_scope() as session:
        return [dict(employee_id=u.employee_id, full_name=u.full_name, role=u.role,
                     email=u.email, contact_number=u.contact_number)
                for u in session.query(User).filter(User.active.is_(True)).order_by(User.full_name)]


def require_admin(user_id):
    sync_directory()
    with session_scope() as session:
        user = session.get(User, user_id)
        if not user or not user.active or user.role not in ADMIN_ROLES:
            raise PermissionError('Only Admin and Director can perform this action.')


def _save_row(path, headers, identity, values, user_id, original_identity=None, employee_rename=None, delete=False):
    """Validate and back up related workbooks before synchronizing SQL in one transaction."""
    with _LOCK:
        require_admin(user_id)
        path = Path(path)
        originals = {path: path.read_bytes()}
        workbooks = {path: load_workbook(path)}
        workbook = workbooks[path]
        sheet, header_row, columns = _sheet(workbook, headers, (identity,))
        lookup = original_identity if original_identity is not None else values[identity]
        target = next((r for r in range(header_row + 1, sheet.max_row + 1)
                       if _text(sheet.cell(r, columns[identity]).value).casefold() == _text(lookup).casefold()), None)
        if target is None and original_identity is not None:
            workbook.close()
            raise DatasetError('The original record no longer exists. Reload and try again.')
        if target is None:
            target = next((r for r in range(header_row + 1, sheet.max_row + 2)
                           if not any(sheet.cell(r, c).value for c in columns.values())), sheet.max_row + 1)
        if identity == 'name' and 'serial' in columns and not sheet.cell(target, columns['serial']).value:
            serials = [sheet.cell(r, columns['serial']).value for r in range(header_row + 1, sheet.max_row + 1)]
            sheet.cell(target, columns['serial']).value = max([v for v in serials if isinstance(v, (int, float))] or [0]) + 1
        for field, value in values.items():
            if field in columns:
                cell = sheet.cell(target, columns[field])
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'
        if delete:
            sheet.delete_rows(target)
        if employee_rename:
            employee_id, new_name = employee_rename
            _, clients = read_directory()
            assigned = {c['name'] for c in clients if c['assigned_employee_id'] == employee_id}
            client_path = Path(CLIENT_FILE)
            originals[client_path] = client_path.read_bytes()
            linked = workbooks[client_path] = load_workbook(client_path)
            client_sheet, first, client_columns = _sheet(linked, CLIENT_HEADERS, ('name', 'assigned_technician'))
            for row in range(first + 1, client_sheet.max_row + 1):
                if _text(client_sheet.cell(row, client_columns['name']).value) in assigned:
                    cell = client_sheet.cell(row, client_columns['assigned_technician'])
                    cell.value, cell.data_type = new_name, 's'
        staged, replaced = {}, []
        try:
            for destination, book in workbooks.items():
                with NamedTemporaryFile(suffix='.xlsx', dir=destination.parent, delete=False) as temp:
                    staged[destination] = Path(temp.name)
                book.save(staged[destination])
            employees, _ = read_directory(staged.get(Path(EMPLOYEE_FILE)), staged.get(Path(CLIENT_FILE)))
            if not any(e['role'] in ADMIN_ROLES for e in employees):
                raise DatasetError('Keep at least one Admin or Director account.')
            if employee_rename and sum(_key(e['full_name']) == _key(employee_rename[1]) for e in employees) != 1:
                raise DatasetError('Employee names must be unambiguous for client assignments.')
            if any(p.read_bytes() != content for p, content in originals.items()):
                raise DatasetError('Workbook changed while editing. Reload and try again.')
            backup_dir = BASE_DIR / 'storage' / 'dataset_backups'
            backup_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
            for destination, content in originals.items():
                backup = backup_dir / f'{destination.stem}-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx'
                backup.write_bytes(content)
                backup.chmod(0o600)
            for destination, temp_path in staged.items():
                os.replace(temp_path, destination)
                replaced.append(destination)
            rename = (original_identity, values['name']) if identity == 'name' and original_identity else None
            sync_directory(force=True, client_rename=rename)
        except Exception:
            for destination in replaced:
                with NamedTemporaryFile(dir=destination.parent, delete=False) as restore:
                    restore.write(originals[destination])
                    restore_path = Path(restore.name)
                os.replace(restore_path, destination)
            raise
        finally:
            for book in workbooks.values():
                book.close()
            for temp_path in staged.values():
                temp_path.unlink(missing_ok=True)


def save_employee(user_id, values):
    require_admin(user_id)
    from services.auth_service import hash_password
    values = dict(values)
    employees, _ = read_directory()
    existing = next((e for e in employees if e['employee_id'] == values.get('employee_id')), None)
    rename = (existing['employee_id'], values['full_name']) if existing and values.get('full_name') != existing['full_name'] else None
    if values.get('password'):
        values['password'] = hash_password(values['password'])
    elif existing:
        values.pop('password', None)
    else:
        raise DatasetError('A password is required for a new employee.')
    _save_row(EMPLOYEE_FILE, EMP_HEADERS, 'employee_id', values, user_id, employee_rename=rename)


def save_client(user_id, values, original_name=None):
    require_admin(user_id)
    values = dict(values)
    identifier = values.pop('assigned_employee_id', '')
    employees = employee_directory()
    match = next((e for e in employees if e['employee_id'] == identifier), None)
    if not match:
        raise DatasetError('Select an assigned employee.')
    values['assigned_technician'] = match['full_name']
    values['office'] = values['office'].replace(' Office', '')
    _save_row(CLIENT_FILE, CLIENT_HEADERS, 'name', values, user_id, original_identity=original_name)


def delete_employee(user_id, employee_id):
    """Remove a directory entry; retain the inactive SQL user for report history."""
    with _LOCK:
        require_admin(user_id)
        with session_scope() as session:
            actor = session.get(User, user_id)
            if actor.employee_id == employee_id:
                raise DatasetError('You cannot delete your own account.')
        employees, clients = read_directory()
        employee = next((e for e in employees if e['employee_id'] == employee_id), None)
        if not employee:
            raise DatasetError('Employee no longer exists. Reload and try again.')
        if not any(e['role'] in ADMIN_ROLES and e['employee_id'] != employee_id for e in employees):
            raise DatasetError('Keep at least one Admin or Director account.')
        if any(c['assigned_employee_id'] == employee_id or
               _key(c['assigned_technician']) == _key(employee['full_name']) for c in clients):
            raise DatasetError('Reassign this employee\'s clients in the Clients tab before deleting.')
        _save_row(EMPLOYEE_FILE, EMP_HEADERS, 'employee_id', {'employee_id': employee_id},
                  user_id, original_identity=employee_id, delete=True)


def delete_client(user_id, name):
    """Remove from the active directory without deleting saved reports."""
    _save_row(CLIENT_FILE, CLIENT_HEADERS, 'name', {'name': name},
              user_id, original_identity=name, delete=True)
