"""Create an ignored deployment secrets file; never print its contents."""
import base64
import io
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from services.auth_service import hash_password
from services.dataset_service import EMPLOYEE_FILE, CLIENT_FILE, EMP_HEADERS, _sheet, read_directory


def main():
    read_directory()  # Validate before preparing a deployment snapshot.
    workbook = load_workbook(EMPLOYEE_FILE)
    sheet, header_row, columns = _sheet(workbook, EMP_HEADERS, EMP_HEADERS.values())
    for row in range(header_row + 1, sheet.max_row + 1):
        if not sheet.cell(row, columns['employee_id']).value:
            continue
        cell = sheet.cell(row, columns['password'])
        password = str(cell.value)
        if not password.startswith('pbkdf2$'):
            cell.value = hash_password(password)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    employee_data = base64.b64encode(stream.getvalue()).decode('ascii')
    client_data = base64.b64encode(Path(CLIENT_FILE).read_bytes()).decode('ascii')
    destination = ROOT / 'storage/deployment/secrets.toml'
    destination.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    destination.write_text(
        '# PRIVATE: paste into Streamlit Cloud Advanced settings > Secrets. Never commit.\n'
        f'EMPLOYEE_DATA_BASE64 = "{employee_data}"\n'
        f'CLIENT_DATA_BASE64 = "{client_data}"\n', encoding='utf-8')
    destination.chmod(0o600)
    print('Prepared storage/deployment/secrets.toml with hashed employee passwords. Source workbooks unchanged.')


if __name__ == '__main__':
    main()
